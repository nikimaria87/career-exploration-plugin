from pathlib import Path
import json, textwrap
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

R=Path(__file__).resolve().parents[1]
P='#4C1734';G='#766B70';I='#241C24';ROSE='#E3B4C6';PAPER='#FBF7F9'
pdfmetrics.registerFont(TTFont('NJSans',str(R/'fonts/DejaVuSans.ttf')))
pdfmetrics.registerFont(TTFont('NJBold',str(R/'fonts/DejaVuSans-Bold.ttf')))
def line(c,x,y,t,size=11,color=I,bold=False):
 c.setFont('NJBold' if bold else 'NJSans',size);c.setFillColor(HexColor(color));c.drawString(x,y,t)
def paragraph(c,x,y,t,width=500,size=11,leading=18,color=G):
 words=t.split();current='';out=[]
 for w in words:
  trial=(current+' '+w).strip()
  if pdfmetrics.stringWidth(trial,'NJSans',size)>width and current:out.append(current);current=w
  else:current=trial
 if current:out.append(current)
 for v in out:line(c,x,y,v,size,color);y-=leading
 return y
def image(c,rel,x,y,w,h):c.drawImage(str(R/rel),x,y,width=w,height=h,mask='auto')
def head(c,n,title):
 image(c,'logos/nextjobber-wordmark-primary.png',48,724,290,47.125)
 line(c,48,672,title,28,P,True);line(c,542,35,f'{n:02}',9,G)

c=canvas.Canvas(str(R/'guidelines/NextJobber_Brand_Guide.pdf'),pagesize=(612,792));c.setTitle('NextJobber Brand Guidelines')
head(c,1,'Brand foundations')
paragraph(c,48,617,'NextJobber helps independent job seekers explore career options and make progress toward their next role.',500,15,24,I)
line(c,48,510,'Your next career step, with a plan.',23,P)
image(c,'icons/nextjobber-profile-plum.png',350,218,200,200)
line(c,48,412,'Original identity',16,I,True)
paragraph(c,48,380,'Airy geometric lettering, complete letter stems, rounded joins, and a custom NJ symbol form the new identity. The supplied master artwork is newly drawn.',250,11,18)
line(c,48,253,'Brand character',16,I,True)
paragraph(c,48,220,'Clear, considered, encouraging. Focus on choices, real examples, and a useful next action.',250,11,18)
line(c,48,104,'Version 1.0  /  October 22, 2026',10,G)
c.showPage();head(c,2,'Logo system')
image(c,'logos/nextjobber-wordmark-primary.png',48,516,516,84)
c.setFillColor(HexColor(P));c.rect(48,334,516,140,fill=1,stroke=0)
image(c,'logos/nextjobber-wordmark-white.png',63,354,486,79)
image(c,'logos/nextjobber-symbol.png',48,160,104,104)
paragraph(c,180,243,'Use the NJ symbol for profile icons, favicons, and small placements. Keep it at least 24 pixels wide on screen.',375,11,18)
paragraph(c,48,114,'Clear space: at least half the capital-letter height on every side. Minimum wordmark: 180 screen pixels or 45 mm in print. Preserve proportions and stroke weights.',516,11,18)
c.showPage();head(c,3,'Color and typography')
for j,(name,color) in enumerate([('Plum',P),('Warm gray',G),('Ink',I),('Rose',ROSE),('Paper',PAPER)]):
 x=48+j*104;c.setFillColor(HexColor(color));c.rect(x,501,92,92,fill=1,stroke=0);line(c,x,478,name,10,I,True);line(c,x,460,color,9,G)
line(c,48,391,'Your experience has more to offer',23,P,True)
line(c,48,344,'DejaVu Sans Bold for headings',13,I,True)
paragraph(c,48,310,'DejaVu Sans for body text. Both font files and the redistribution license are included. The logo is outlined artwork and does not need a font installation.',510,12,20)
line(c,48,221,'Readable pairings',16,I,True)
paragraph(c,48,188,'Use ink or plum on light backgrounds and white on plum. Reserve rose for accents rather than small text on white. Preserve strong contrast in buttons, captions, and print templates.',510,12,20)
c.showPage();head(c,4,'Asset library')
items=json.loads((R/'asset-categories.json').read_text());y=615
for item in items:
 line(c,48,y,item['category'],11,I,True);line(c,230,y,item['description'][:56],9,G);y-=25
paragraph(c,48,140,'Replace bracketed contact and invoice fields before use. The QR opens the current private dashboard. Review its destination and access before printing publicly.',516,10,16)
paragraph(c,48,77,'The support widget is a local prototype and response library. Generated mockups illustrate the identity; they are not photographs of a live deployment.',516,9,14)
c.save()

c=canvas.Canvas(str(R/'logos/NextJobber_Logos_Print.pdf'),pagesize=(960,300));c.setTitle('NextJobber logo print exports')
for file,bg in [('nextjobber-wordmark-primary.png','#FFFFFF'),('nextjobber-wordmark-plum.png','#FFFFFF'),('nextjobber-wordmark-white.png',P)]:
 c.setFillColor(HexColor(bg));c.rect(0,0,960,300,fill=1,stroke=0);image(c,'logos/'+file,48,80,864,140.4);c.showPage()
c.save()
for face in ['front','back']:
 c=canvas.Canvas(str(R/f'print/business-card-{face}.pdf'),pagesize=(270,162));c.setTitle('NextJobber business card '+face+' with bleed');image(c,f'print/business-card-{face}.png',0,0,270,162);c.save()
c=canvas.Canvas(str(R/'qr/NextJobber_QR_Print.pdf'),pagesize=(432,576));c.setTitle('NextJobber dashboard QR')
image(c,'logos/nextjobber-wordmark-primary.png',36,473,360,58.5);image(c,'qr/nextjobber-dashboard-qr.png',72,148,288,288);line(c,72,126,'Open the current NextJobber workspace',12,P)
paragraph(c,36,84,'Current destination requires owner access. Review the URL and audience before public printing.',360,10,16);c.save()

d=Document();s=d.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11);s.top_margin=s.bottom_margin=Inches(.7);s.left_margin=s.right_margin=Inches(.75)
for name in ['Normal','Title','Heading 1','Heading 2']:
 st=d.styles[name];st.font.name='DejaVu Sans';st.font.color.rgb=RGBColor(0,0,0);st.font.size=Pt(10 if name=='Normal' else 24)
d.styles['Normal'].paragraph_format.space_after=Pt(8)
d.add_picture(str(R/'logos/nextjobber-wordmark-primary.png'),width=Inches(3.0))
d.add_paragraph('Invoice template','Title')
d.add_paragraph('[Legal business name]\n[Business address]\n[Email or phone]')
d.add_paragraph('Invoice number  [INV-0001]     Date  [YYYY-MM-DD]\nPayment due  [YYYY-MM-DD]     Currency  [USD]')
d.add_paragraph('Bill to','Heading 2');d.add_paragraph('[Client name]\n[Client address]\n[Client email]')
t=d.add_table(rows=1,cols=4);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
widths=[3.8,.7,1.1,1.4]
for column,w in zip(t.columns,widths):column.width=Inches(w)
for cell,w,title in zip(t.rows[0].cells,widths,['Description','Qty','Rate','Amount']):cell.width=Inches(w);cell.text=title
for desc in ['[Service description]','[Additional service]','[Additional service]']:
 row=t.add_row()
 for cell,w,value in zip(row.cells,widths,[desc,'[1]','[0.00]','[0.00]']):cell.width=Inches(w);cell.text=value
for ri,row in enumerate(t.rows):
 for ci,cell in enumerate(row.cells):
  cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  tcPr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
  for edge in ['top','left','bottom','right']:
   v=OxmlElement('w:'+edge);v.set(qn('w:val'),'single');v.set(qn('w:sz'),'4');v.set(qn('w:color'),'D9D9D9');borders.append(v)
  tcPr.append(borders);m=OxmlElement('w:tcMar')
  for edge in ['top','left','bottom','right']:
   v=OxmlElement('w:'+edge);v.set(qn('w:w'),'100');v.set(qn('w:type'),'dxa');m.append(v)
  tcPr.append(m)
  if ri==0:
   shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'3D3440');tcPr.append(shade)
  for p in cell.paragraphs:
   p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
   p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.RIGHT
   for run in p.runs:run.font.size=Pt(9);run.font.bold=ri==0;run.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '241C24')
p=d.add_paragraph('Subtotal  [0.00]\nTax, if applicable  [0.00]\nTotal due  [0.00]');p.paragraph_format.space_before=Pt(12)
d.add_paragraph('Payment instructions','Heading 2');d.add_paragraph('[Approved payment method and instructions]')
p=d.add_paragraph('Replace bracketed fields and calculate totals before sending.');p.runs[0].font.size=Pt(8)
for root in [d._element,d.styles._element]:
 for border in list(root.iter(qn('w:pBdr'))):border.getparent().remove(border)
d.core_properties.author='NextJobber';d.core_properties.title='NextJobber Invoice Template';d.save(R/'print/NextJobber_Invoice_Template.docx')
print('Brand guide, print PDFs, QR PDF, and editable invoice created.')
